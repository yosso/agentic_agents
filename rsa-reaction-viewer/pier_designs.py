"""Prepare traceable workbook inputs. No OpenSees model is run here."""
from dataclasses import dataclass, field, asdict
from pathlib import Path
import argparse
import json
import math
import re
import openpyxl
from openpyxl.utils.cell import range_boundaries
import config


@dataclass
class Issue:
    severity: str
    footing: str
    source: str
    message: str


@dataclass
class SoilLayer:
    source_row: int
    properties: dict
    cells: dict


@dataclass
class SoilProfile:
    boring_id: str
    table_name: str
    id_source: str
    layers: list[SoilLayer] = field(default_factory=list)
    units: dict = field(default_factory=dict)


@dataclass
class PierDesign:
    footing: str
    boring_id: str
    soil_table: str | None
    diameter_ft: float | None
    minimum_embedment_ft: float | None
    water_table_depth_ft: float | None
    allowable_deflection_in: float | None
    reactions: list[dict]
    source_cells: dict


@dataclass
class Preparation:
    workbook: str
    soil_profiles: list[SoilProfile] = field(default_factory=list)
    piers: list[PierDesign] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)

    @property
    def input_valid(self):
        return bool(self.piers) and not any(i.severity == 'ERROR' for i in self.issues)

    def to_dict(self):
        return dict(asdict(self), input_valid=self.input_valid, solver_ready=False,
                    stage='Workbook preparation only; soil models, section materials, load conventions and solver settings still require validation.',
                    reaction_units={'force': 'N', 'moment': 'N-m'})


def normalize(value):
    return ' '.join(str(value or '').split()).upper()


def boring_key(value):
    text = normalize(value)
    match = re.fullmatch(r'(?:SB\s*)?0*(\d+)(?:\.0)?', text)
    return f'SB{int(match[1])}' if match else text


SOIL_FIELDS = {
    'START': 'start', 'END': 'end', 'MODEL': 'model', 'DENSITY': 'density',
    'COHESION ULT.(KSF)': 'cohesion_ultimate', 'COHESION FACT.': 'cohesion_factored',
    'FRICTION ANGLE': 'friction_angle', 'MODULUS OF SUBGRADE REACTION (STATIC)': 'k_static',
    'MODULUS OF SUBGRADE REACTION (CYCLIC)': 'k_cyclic', 'FKH': 'k_factored',
    'E50': 'e50', 'FE50': 'e50_factored',
}
ASSIGNMENT_FIELDS = {
    'FTG MARK': 'footing', 'BORING': 'boring', 'PIER DIA (FT)': 'diameter',
    'MIN EMBED DEPTH (FT)': 'minimum_depth', 'WATER TABLE DEPTH (FT)': 'water_depth',
    'ALLOWABLE DEFLECTION (IN)': 'allowable',
}


def prepare_pier_designs(workbook_path, reactions=(), *, boring_ids=None, soil_units=None):
    """Join each footing to a boring profile and complete node/combination rows.

    Boring IDs come from the numeric table suffix. The NUMBER/BORING cell and
    optional configured IDs are cross-checks, not overrides. Soil units are explicit metadata.
    Errors are collected so all missing/ambiguous inputs can be reviewed together.
    """
    overrides = config.SOIL_BORING_IDS if boring_ids is None else boring_ids
    units = dict(config.SOIL_UNITS if soil_units is None else soil_units)
    path = Path(workbook_path).resolve()
    result = Preparation(str(path))
    values = openpyxl.load_workbook(path, data_only=True, keep_links=False)
    formulas = None
    def issue(source, message, footing='', severity='ERROR'):
        result.issues.append(Issue(severity, footing, source, message))
    try:
        formulas = openpyxl.load_workbook(path, data_only=False, keep_links=False)
        formula_cells = []
        def read(sheet, row, column):
            cell = values[sheet].cell(row, column)
            formula = formulas[sheet].cell(row, column)
            source = f'{sheet}!{cell.coordinate}'
            if cell.data_type == 'e':
                issue(source, f'Excel error: {cell.value}')
                return None
            if formula.data_type == 'f':
                formula_cells.append(source)
                if cell.value is None:
                    issue(source, 'Formula has no saved result. Recalculate and save in Excel.')
            return cell.value

        def table_rows(sheet_name, table_name, fields):
            sheet = values[sheet_name]
            table = sheet.tables[table_name]
            left, top, right, bottom = range_boundaries(table.ref)
            columns = {}
            for col in range(left, right + 1):
                header = normalize(sheet.cell(top, col).value)
                if header in fields:
                    name = fields[header]
                    if name in columns:
                        issue(f'{sheet_name}/{table_name}', f'Duplicate field: {header}')
                    columns[name] = col
            missing = set(fields.values()) - set(columns)
            if missing:
                issue(f'{sheet_name}/{table_name}', 'Missing columns: ' + ', '.join(sorted(missing)))
            for row in range(top + 1, bottom + 1 - (table.totalsRowCount or 0)):
                data = {name: read(sheet_name, row, col) for name, col in columns.items()}
                if all(v is None or str(v).strip() == '' for v in data.values()):
                    continue
                cells = {name: f'{sheet_name}!{sheet.cell(row, col).coordinate}' for name, col in columns.items()}
                yield row, data, cells

        def number(value, source, footing='', positive=False, optional=False):
            if value is None or str(value).strip() in ('', '-'):
                if not optional:
                    issue(source, 'Required numeric input is missing.', footing)
                return None
            try:
                if isinstance(value, bool):
                    raise ValueError
                value = float(value)
                if not math.isfinite(value) or value < 0 or (positive and value == 0):
                    raise ValueError
                return value
            except (TypeError, ValueError):
                issue(source, 'Expected a finite ' + ('positive' if positive else 'nonnegative') + ' number.', footing)
                return None

        soil_sheet = config.SOIL_SHEET_NAME
        if soil_sheet not in values:
            issue(soil_sheet, 'Soil worksheet not found.')
        else:
            sheet = values[soil_sheet]
            tables = [t for t in sheet.tables.values() if t.name.startswith(config.SOIL_TABLE_PREFIX)]
            if not tables:
                issue(soil_sheet, f'No tables starting with {config.SOIL_TABLE_PREFIX}.')
            for unknown in set(overrides) - {t.name for t in tables}:
                issue(soil_sheet, f'Boring-ID override refers to missing table {unknown}.')
            for table in tables:
                left, top, _, _ = range_boundaries(table.ref)
                id_cell = f'{soil_sheet}!{sheet.cell(top - 1, left + 1).coordinate}' if top > 1 else table.name
                suffix = table.name[len(config.SOIL_TABLE_PREFIX):]
                if not re.fullmatch(r'\d+', suffix):
                    issue(table.name, 'Expected a numeric boring suffix, such as _03 for boring 3.')
                    continue
                boring = f'SB{int(suffix)}'
                if table.name in overrides and boring_key(overrides[table.name]) != boring:
                    issue(table.name, f'Configured boring ID {overrides[table.name]!r} disagrees with table boring {boring}.')
                if top > 1 and normalize(sheet.cell(top - 1, left).value) in ('NUMBER', 'BORING'):
                    label = read(soil_sheet, top - 1, left + 1)
                    if label is not None and str(label).strip() and boring_key(label) != boring:
                        issue(id_cell, f'Boring ID mismatch: table {table.name} identifies {boring}, but worksheet label is {label!r}.')
                profile = SoilProfile(boring, table.name, f'{soil_sheet}/{table.name}', units=units.copy())
                previous_end = None
                for row, data, cells in table_rows(soil_sheet, table.name, SOIL_FIELDS):
                    properties = {}
                    for name, value in data.items():
                        properties[name] = (str(value).strip() if value is not None else '') if name == 'model' else number(value, cells[name], optional=name == 'k_cyclic')
                    if not properties.get('model'):
                        issue(cells.get('model', table.name), 'Soil model name is missing.')
                    start, end = properties.get('start'), properties.get('end')
                    if start is not None and end is not None:
                        if end <= start:
                            issue(cells['end'], 'Layer END must exceed START.')
                        expected_start = previous_end if previous_end is not None else 0
                        if not math.isclose(start, expected_start, abs_tol=1e-8):
                            issue(cells['start'], f'Layer gap/overlap or missing surface layer: expected START {expected_start}, got {start}.')
                        previous_end = end
                    if properties.get('friction_angle') is not None and properties['friction_angle'] >= 90:
                        issue(cells['friction_angle'], 'Friction angle must be less than 90 degrees.')
                    profile.layers.append(SoilLayer(row, properties, cells))
                if not profile.layers:
                    issue(table.name, 'Soil profile has no populated layers.')
                result.soil_profiles.append(profile)

        profiles = {}
        for profile in result.soil_profiles:
            key = boring_key(profile.boring_id)
            profiles.setdefault(key, []).append(profile)
        for key, matches in profiles.items():
            if len(matches) > 1:
                issue(soil_sheet, f'Ambiguous boring {key}: ' + ', '.join(p.table_name for p in matches))
        required_units = ['depth', 'density', 'cohesion', 'subgrade_modulus', 'strain', 'friction_angle']
        for key in required_units:
            if not units.get(key):
                issue(soil_sheet, f'Confirm soil unit/convention for {key} in SOIL_UNITS; no conversion has been assumed.')
        issue(soil_sheet, 'Cohesion header states KSF; confirm that values and selected unit agree. Confirm ultimate/factored parameter selection before building soil springs.', severity='WARNING')

        assignment_sheet, assignment_table = config.ASSIGNMENT_SHEET_NAME, config.ASSIGNMENT_TABLE_NAME
        if assignment_sheet not in values or assignment_table not in values[assignment_sheet].tables:
            issue(assignment_table, 'Footing-to-boring assignment table not found.')
        else:
            seen = set()
            for _, data, cells in table_rows(assignment_sheet, assignment_table, ASSIGNMENT_FIELDS):
                mark = str(data.get('footing') or '').strip()
                boring = str(data.get('boring') or '').strip()
                if not mark:
                    issue(cells.get('footing', assignment_table), 'Footing mark is missing.')
                    continue
                if normalize(mark) in seen:
                    issue(cells['footing'], 'Duplicate footing assignment.', mark)
                seen.add(normalize(mark))
                matches = profiles.get(boring_key(boring), [])
                profile = matches[0] if len(matches) == 1 else None
                if profile is None:
                    issue(cells.get('boring', assignment_table), f'Boring {boring!r} has no unique soil table.', mark)
                loads = [dict(r) for r in reactions if normalize(r.get('mark')) == normalize(mark)]
                if not loads:
                    issue(assignment_table, 'No extracted load combinations for this footing.', mark)
                for load in loads:
                    for name in ['model', 'group', 'node', 'case', 'case_name']:
                        if load.get(name) is None or str(load[name]).strip() == '':
                            issue('RSA reactions', f'Missing source field {name}.', mark)
                    for name in ['FX', 'FY', 'FZ', 'MX', 'MY', 'MZ']:
                        value = load.get(name)
                        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                            issue('RSA reactions', f'Invalid {name} at node {load.get("node")}, combination {load.get("case")}.', mark)
                numeric = {name: number(data.get(name), cells.get(name, assignment_table), mark,
                                       positive=name in ('diameter', 'minimum_depth', 'allowable'))
                           for name in ['diameter', 'minimum_depth', 'water_depth', 'allowable']}
                result.piers.append(PierDesign(mark, boring, profile.table_name if profile else None,
                    numeric['diameter'], numeric['minimum_depth'], numeric['water_depth'], numeric['allowable'], loads, cells))
            for load in reactions:
                if normalize(load.get('mark')) not in seen:
                    issue(assignment_table, 'Extracted footing has no design assignment.', str(load.get('mark', '')))
        if not result.piers:
            issue(assignment_table, 'No pier designs could be assembled.')
        if formula_cells:
            issue('Excel saved formula results', f'Used {len(formula_cells)} cached formula values. Recalculate/save Excel before preparation; this reader does not evaluate VBA/formulas.', severity='WARNING')
        issue(assignment_table, 'Allowable deflection is the assignment-table value pending confirmation of the diameter-table column and override precedence.', severity='WARNING')
        return result
    finally:
        values.close()
        if formulas is not None:
            formulas.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', default=str(config.DEFAULT_WORKBOOK))
    parser.add_argument('--reactions', required=True, help='JSON output of read_robot.ps1, in N and N-m')
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    extraction = json.loads(Path(args.reactions).read_text(encoding='utf-8-sig'))
    report = prepare_pier_designs(args.workbook, extraction.get('reactions', []))
    if extraction.get('error'):
        report.issues.append(Issue('ERROR', '', 'Robot extraction', extraction['error']))
    for group in extraction.get('groups', []):
        if group.get('status') != 'OK':
            report.issues.append(Issue('ERROR', group.get('mark', ''), 'Robot group', group.get('status', 'Not checked')))
    Path(args.output).write_text(json.dumps(report.to_dict(), indent=2, allow_nan=False), encoding='utf-8')
    print(f'{len(report.piers)} piers, {len(report.soil_profiles)} soil profiles, {sum(i.severity == "ERROR" for i in report.issues)} errors. Solver not run.')

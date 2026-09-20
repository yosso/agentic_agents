"""Workbook input and reaction presentation. Never saves the workbook."""
from pathlib import Path
import math
import re
import openpyxl
from openpyxl.utils.cell import range_boundaries
import config

def read_mapping(path, sheet_name, table_name):
    path = Path(path).resolve()
    book = openpyxl.load_workbook(path, data_only=True, keep_links=False)
    try:
        if sheet_name not in book:
            raise ValueError(f'Sheet {sheet_name!r} not found. Available: {", ".join(book.sheetnames)}')
        sheet = book[sheet_name]
        if table_name not in sheet.tables:
            raise ValueError(f'Table {table_name!r} not found. Available: {", ".join(sheet.tables)}')
        table = sheet.tables[table_name]
        left, top, right, bottom = range_boundaries(table.ref)
        headers = [str(sheet.cell(top, col).value or '').strip() for col in range(left, right + 1)]
        required = [config.MARK_COLUMN, config.QUANTITY_COLUMN, config.GROUP_COLUMN, config.MODEL_COLUMN]
        if len(set(headers)) != len(headers):
            raise ValueError('Mapping table contains duplicate column headers.')
        missing = set(required) - set(headers)
        if missing:
            raise ValueError(f'Missing mapping columns: {", ".join(sorted(missing))}')
        mappings = []
        seen = set()
        for row in range(top + 1, bottom + 1 - (table.totalsRowCount or 0)):
            values = {h: sheet.cell(row, left + i).value for i, h in enumerate(headers)}
            if all(v is None or str(v).strip() == '' for v in values.values()):
                continue
            if any(values[h] is None or str(values[h]).strip() == '' for h in required):
                raise ValueError(f'Row {row}: fill all four mapping fields.')
            mark, group, filename = [str(values[h]).strip() for h in
                                     [config.MARK_COLUMN, config.GROUP_COLUMN, config.MODEL_COLUMN]]
            quantity = float(values[config.QUANTITY_COLUMN])
            if not math.isfinite(quantity) or quantity < 1 or not quantity.is_integer():
                raise ValueError(f'Row {row}: footing quantity must be a positive whole number.')
            model = Path(filename)
            if not model.suffix:
                model = model.with_suffix('.rtd')
            if not model.is_absolute():
                model = path.parent / model
            model = model.resolve()
            key = (str(model).casefold(), group.casefold())
            if key in seen:
                raise ValueError(f'Row {row}: model/group appears more than once. Resolve the duplicate assignment.')
            seen.add(key)
            mappings.append(dict(mark=mark, expected=int(quantity), group=group,
                                 model=str(model), row=row, exists=model.is_file(),
                                 source=f'{path.name} / {sheet_name} / {table_name}',
                                 quantity_cell=sheet.cell(row, left + headers.index(config.QUANTITY_COLUMN)).coordinate))
        if not mappings:
            raise ValueError('The mapping table has no populated rows.')
        return mappings
    finally:
        book.close()

def parse_cases(text):
    if not text.strip():
        return []
    cases = set()
    for token in re.split(r'[,;\s]+', text.strip()):
        if not re.fullmatch(r'\d+(?:-\d+)?', token):
            raise ValueError('Enter combination numbers such as 1, 2, 5-10, or leave blank for all supported combinations.')
        ends = [int(n) for n in token.split('-')]
        start, end = ends[0], ends[-1]
        if start < 1 or end < start or end - start > 10000:
            raise ValueError('Invalid combination range.')
        cases.update(range(start, end + 1))
    return sorted(cases)

def display_reaction(row, units):
    force, moment = {'kip / kip-ft': (4448.2216152605, 1355.8179483314),
                     'kN / kN-m': (1000, 1000), 'N / N-m': (1, 1)}[units]
    return [row[k] / (force if k.startswith('F') else moment)
            for k in ['FX', 'FY', 'FZ', 'MX', 'MY', 'MZ']]


def assess_group(group, mapping):
    """Compare only measured counts; an unavailable read is not a zero count."""
    expected = mapping['expected']
    source = f"{mapping.get('source', 'Excel row ' + str(mapping['row']))}!{mapping.get('quantity_cell', '')}"
    findings = []
    nodes, supports = group.get('node_count'), group.get('support_count')
    if nodes is not None and nodes != expected:
        findings.append(f'FOOTING QTY mismatch at {source}: Excel {expected}, Robot group nodes {nodes} (difference {nodes - expected:+d}).')
    if supports is not None and supports != expected and supports != nodes:
        findings.append(f'Support quantity mismatch at {source}: Excel {expected}, Robot supports {supports}.')
    status = group.get('status', '')
    if status and status not in ('OK', 'Quantity mismatch'):
        findings.append(status)
    if nodes is None and not findings:
        findings.append('Robot quantity could not be verified.')
    if status == 'Quantity mismatch' and not findings:
        findings.append(status)
    return ('REVIEW' if findings else 'MATCH'), findings, source


def controlling_combinations(reactions):
    """Five governing checks per footing callout; loads are never summed."""
    footings = {}
    for row in reactions:
        footings.setdefault(row['mark'], []).append(row)
    output = []
    for mark, rows in sorted(footings.items()):
        rows = sorted(rows, key=lambda r: (r['model'], r['group'], r['node'], r['case']))
        criteria = [('Maximum shear', 'H'), ('Maximum bending moment', 'MH'),
                    ('Maximum torsion', 'MZ'), ('Maximum compression', 'C'),
                    ('Maximum uplift', 'U')]
        for label, component in criteria:
            eligible = [r for r in rows if (component != 'C' or r['FZ'] > 0)
                        and (component != 'U' or r['FZ'] < 0)]
            def value(row):
                if component == 'H':
                    return math.hypot(row['FX'], row['FY'])
                if component == 'MH':
                    return math.hypot(row['MX'], row['MY'])
                if component == 'MZ':
                    return abs(row['MZ'])
                return row['FZ']
            extreme = ((min if component == 'U' else max)(value(r) for r in eligible)
                       if eligible else None)
            tied = [r for r in eligible if value(r) == extreme]
            output.append(dict(mark=mark, criterion=label, value=extreme,
                               component=component, ties=len(tied),
                               reaction=tied[0] if tied else None))
    return output

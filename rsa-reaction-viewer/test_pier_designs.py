import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from pier_designs import prepare_pier_designs, SOIL_FIELDS, ASSIGNMENT_FIELDS
from preparation_report import soil_report_rows, markdown_report


class PreparationTests(unittest.TestCase):
    def run_fixture(self, boring='SB3', end=10, reactions=None, assignment='SB3', descriptive_header=False, soil_units=None, start=0, next_start=None, percentage_cells=False, multiline_cohesion=False):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'inputs.xlsx'
            book = Workbook()
            soil = book.active
            soil.title = 'GEOTECH PARAMETERS'
            soil.append(['NUMBER', boring])
            headers = list(SOIL_FIELDS)
            if multiline_cohesion:
                headers[5] = 'COHESION\nFACT.(PSF)'
            row = [start, end, 'SOFT CLAY', 125, 1, 0.66, 0, 200, '-', 133, 1, 1.5]
            if descriptive_header:
                headers[9] = 'FACTORED MODULUS OF SUBGRADE REACTION (STATIC)'
            soil.append(headers)
            soil.append(row)
            if percentage_cells:
                soil['K3'] = 0.01
                soil['L3'] = 0.015
                soil['K3'].number_format = '0.0%'
                soil['L3'].number_format = '0.0%'
            if next_start is not None:
                soil.append([next_start, 20, *row[2:]])
            soil.add_table(Table(displayName='tblSOIL_BORING_03', ref='A2:L4' if next_start is not None else 'A2:L3'))
            sheet = book.create_sheet('INPUT TABLES')
            sheet.append(list(ASSIGNMENT_FIELDS))
            sheet.append(['A', assignment, 2.5, 5, 0, 0.3])
            sheet.add_table(Table(displayName='tblBoringAssignment', ref='A1:F2'))
            criteria = book.create_sheet('DP WORKFLOW and GUIDELINES')
            criteria.append(['DIA (FT)', 'TIP DEF (IN)', 'NON-RECOVERABLE (IN)', 'MIN. EMBED DEPTH'])
            criteria.append([2.5, 0.9, 0.297, 10])
            criteria.add_table(Table(displayName='tblDEFLECTION_SIZE', ref='A1:D2'))
            book.save(path)
            book.close()
            before = path.read_bytes()
            result = prepare_pier_designs(path, reactions or [], soil_units={} if soil_units is None else soil_units)
            self.assertEqual(before, path.read_bytes())
            return result

    def test_id_comes_from_table_suffix(self):
        row = dict(mark='A', model='file', group='g', node=1, case=20, case_name='comb',
                   FX=1, FY=2, FZ=-3, MX=4, MY=5, MZ=6)
        result = self.run_fixture(reactions=[row])
        self.assertEqual(result.soil_profiles[0].boring_id, 'SB3')
        self.assertEqual(result.piers[0].soil_table, 'tblSOIL_BORING_03')
        self.assertEqual(result.piers[0].reactions[0]['FZ'], -3)
        self.assertEqual(result.piers[0].water_table_depth_ft, 0)
        self.assertIsNone(result.soil_profiles[0].layers[0].properties.k_cyclic)
        self.assertFalse(result.input_valid)  # units still unconfirmed
        self.assertFalse(result.to_dict()['solver_ready'])

    def test_unknown_boring_and_invalid_layer_report_sources(self):
        result = self.run_fixture(boring='SB3', end=0, assignment='SB23')
        messages = [i.message for i in result.issues]
        self.assertTrue(any('no unique soil table' in m for m in messages))
        self.assertTrue(any('END depth must exceed' in m for m in messages))
        self.assertTrue(any('No extracted load combinations' in m for m in messages))

    def test_label_mismatch_is_flagged_and_does_not_reassign_profile(self):
        result = self.run_fixture(boring='SB23')
        self.assertEqual(result.soil_profiles[0].boring_id, 'SB3')
        self.assertTrue(any('Boring ID mismatch' in i.message for i in result.issues))

    def test_numeric_label_and_zero_padded_assignment_match(self):
        result = self.run_fixture(boring=3, assignment='SB03')
        self.assertEqual(result.piers[0].soil_table, 'tblSOIL_BORING_03')
        self.assertFalse(any('Boring ID mismatch' in i.message or 'no unique soil table' in i.message for i in result.issues))

    def test_report_preserves_source_missing_values_and_findings(self):
        result = self.run_fixture()
        rows = soil_report_rows(result)
        self.assertEqual(len(rows), 12)
        cyclic = next(r for r in rows if r[4] == 'k_cyclic')
        self.assertEqual(cyclic[5], '-')
        self.assertEqual(cyclic[6], 'Unavailable')
        self.assertEqual(cyclic[8], 'GEOTECH PARAMETERS!I3')
        self.assertIn('UNCONFIRMED', cyclic)
        report = markdown_report(result)
        self.assertIn('tblSOIL_BORING_03', report)
        self.assertIn('INPUT TABLES!C2', report)
        self.assertIn('No extracted load combinations', report)

    def test_factored_static_header_mapping(self):
        for descriptive in (False, True):
            result = self.run_fixture(descriptive_header=descriptive)
            layer = result.soil_profiles[0].layers[0]
            self.assertEqual(layer.properties.k_static_factored, 133)
            self.assertEqual(layer.cells['k_static_factored'], 'GEOTECH PARAMETERS!J3')
            self.assertFalse(hasattr(layer.properties, 'k_cyclic_factored'))
            self.assertFalse(any('Missing columns:' in i.message or 'Ambiguous legacy' in i.message for i in result.issues))

    def test_confirmed_units_and_depths(self):
        import config
        result = self.run_fixture(soil_units=config.SOIL_UNITS)
        layer = result.soil_profiles[0].layers[0]
        self.assertEqual(layer.properties.end, 10)
        self.assertEqual(layer.properties.e50, 0.01)
        self.assertEqual(layer.properties.e50_factored, 0.015)
        self.assertEqual(layer.raw_properties['e50'], 1)
        self.assertEqual(layer.raw_properties['e50_factored'], 1.5)
        self.assertEqual(layer.properties.k_static_factored, 133)
        self.assertFalse(any('Confirm soil unit' in i.message or 'END depth' in i.message for i in result.issues))
        self.assertEqual(result.soil_profiles[0].units['cohesion'], 'lb/ft^2')
        self.assertEqual(layer.properties.cohesion_ultimate, 1)
        self.assertEqual(layer.raw_properties['cohesion_ultimate'], 1)
        rows = soil_report_rows(result)
        self.assertEqual(next(r[7] for r in rows if r[4] == 'density'), 'lb/ft^3')
        self.assertEqual(next(r[7] for r in rows if r[4] == 'cohesion_ultimate'), 'lb/ft^2')

    def test_depth_origin_and_layer_continuity(self):
        valid = self.run_fixture(next_start=10)
        self.assertFalse(any('START depth' in i.message or 'gap/overlap' in i.message for i in valid.issues))
        for start in (9, 11):
            result = self.run_fixture(next_start=start)
            self.assertTrue(any('gap/overlap' in i.message and i.source.endswith('!A4') for i in result.issues))
        result = self.run_fixture(start=1)
        self.assertTrue(any('must be zero' in i.message for i in result.issues))
        result = self.run_fixture(start=-1)
        self.assertTrue(any('nonnegative' in i.message and i.source.endswith('!A3') for i in result.issues))

    def test_percentage_cells_are_not_converted_twice(self):
        layer = self.run_fixture(percentage_cells=True).soil_profiles[0].layers[0]
        self.assertEqual(layer.properties.e50, 0.01)
        self.assertEqual(layer.properties.e50_factored, 0.015)

    def test_psf_cohesion_headers_with_excel_line_break(self):
        result = self.run_fixture(multiline_cohesion=True)
        layer = result.soil_profiles[0].layers[0]
        self.assertEqual(layer.properties.cohesion_ultimate, 1)
        self.assertEqual(layer.properties.cohesion_factored, 0.66)
        self.assertEqual(layer.cells['cohesion_ultimate'], 'GEOTECH PARAMETERS!E3')
        self.assertEqual(layer.cells['cohesion_factored'], 'GEOTECH PARAMETERS!F3')
        self.assertFalse(any('Missing columns' in i.message for i in result.issues))

    def test_typed_parameters_preserve_json_and_source_tracking(self):
        import json
        from soil_parameters import SoilParameters
        result = self.run_fixture()
        layer = result.soil_profiles[0].layers[0]
        self.assertIsInstance(layer.properties, SoilParameters)
        self.assertEqual(layer.properties.model, 'SOFT CLAY')
        self.assertEqual(layer.raw_properties['k_cyclic'], '-')
        exported = json.loads(json.dumps(result.to_dict()))['soil_profiles'][0]['layers'][0]
        self.assertEqual(len(exported['properties']), 12)
        self.assertIsNone(exported['properties']['k_cyclic'])
        self.assertEqual(exported['properties']['k_static_factored'], 133)
        self.assertEqual(exported['cells']['k_static_factored'], 'GEOTECH PARAMETERS!J3')
        self.assertEqual(len(soil_report_rows(result)), 12)

    def test_diameter_table_lookup_preserves_both_limits(self):
        result = self.run_fixture()
        pier = result.piers[0]
        self.assertEqual(pier.minimum_embedment_ft, 10)
        self.assertEqual(pier.diameter_criteria['tip_deflection'], 0.9)
        self.assertEqual(pier.diameter_criteria['nonrecoverable_deflection'], 0.297)
        self.assertEqual(pier.allowable_deflection_in, 0.297)
        self.assertEqual(pier.source_cells['allowable'], 'DP WORKFLOW and GUIDELINES!C2')
        self.assertFalse(any('Select TIP DEF' in i.message for i in result.issues))
        self.assertEqual(pier.source_cells['minimum_depth'], 'DP WORKFLOW and GUIDELINES!D2')
        self.assertTrue(any('Assignment 5.0 differs' in i.message for i in result.issues))


if __name__ == '__main__':
    unittest.main()
